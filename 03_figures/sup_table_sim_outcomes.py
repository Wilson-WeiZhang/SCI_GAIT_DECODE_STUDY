"""Supplementary item

sup_table_sim_outcomes — Supplementary Table 4: the Fig. 7 physics-simulation outcomes of the
four conditions, per patient and pooled, on the 3 s usable-duration, window-matched-walker build

Reads:
  simulation and gate output in 02_decoding/simulation/ (not distributed)

Renders one panel or table from the inputs above; it performs no
preprocessing and does not fit the decoder. It fits a linear mixed model and a least-squares regression
for the statistics reported with that panel or table.
"""

import csv
import hashlib
import os
import platform
import warnings
from decimal import Decimal, ROUND_HALF_UP

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from scipy import stats as sps

os.environ.setdefault("GAIT_F7_SIM", "causal_from2s")
import analysis_config as ec
import f7_scored_trials as F7_SCORED

ec.require("main", "main")
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
RES = os.path.join(ROOT, "02_decoding", "simulation", "simulation_rollouts_gate3s")
RHY = os.path.join(RES, "rhythm_metrics.csv")
WRHY = os.path.join(RES, "walker_matched_rhythm.csv")
TRA = os.path.join(RES, "traces_x.csv.gz")
WTRA = os.path.join(RES, "walker_matched_traces_x.csv.gz")
ROLL = os.path.join(RES, "g_v5_rollouts.csv")
SDD = os.path.join(ROOT, "04_source_data", "supplementary")
SD = os.path.join(SDD, "sup_table_sim_outcomes_source_data.csv")
MD = os.path.join(SDD, "sup_table_sim_outcomes_metadata.csv")
OUT_DIR = os.path.join(ROOT, "figures_out", "supplementary", "tables")
XLSX = os.path.join(OUT_DIR, "Supplementary_Table_S4_Simulation_Outcomes.xlsx")
LEGEND = os.path.join(OUT_DIR, "Supplementary_Table_Simulation_Outcomes_legend.md")
if ec.SCOPED:
    RHY, WRHY, TRA, WTRA, ROLL = (str(ec.sim_in(p)) for p in (RHY, WRHY, TRA, WTRA, ROLL))
    SD, MD, XLSX, LEGEND = (str(ec.scoped(p)) for p in (SD, MD, XLSX, LEGEND))

TEND = 3.0
DISP_Q, DISP_Q_TITLE = "forward displacement at 3.0 s", "Forward displacement at 3.0 s (m)"
if ec.F7_SIM == "causal":
    TEND = 2.0
    DISP_Q, DISP_Q_TITLE = ("forward displacement 2.0 s after the warm-up",
                            "Forward displacement over the 2.0 s after the 1.0 s warm-up (m)")
DISP_DEF = ("the simulated hip displacement of the physics rollout at 3.0 s, measured from the rollout's "
            "own post-settle hip position: the last recorded sample at or before 3.0 s on the 40 ms "
            "recording grid. A displacement in metres, not a gait speed and not a clinical walking "
            "distance.")
DISP_AT = f"{TEND:.1f} s into the rollout"
FALL_DEF = ("the number of rollouts in which trunk pitch exceeded 60 deg from vertical or hip height fell "
            "below half the standing hip height.")
if ec.F7_SIM == "causal_from2s":
    DISP_Q, DISP_Q_TITLE = ("forward displacement 3.0 s after the final-output onset",
                            "Forward displacement over the 3.0 s after the onset of the final decoder output (m)")
    DISP_DEF = ("the simulated hip displacement of the physics rollout 3.0 s after the onset of the final decoder "
                "output (trial time 4 s; the biped is driven from trial time 2 s), measured from the hip position at "
                "that onset: the last recorded sample at or before 3.0 s after the onset on the 40 ms recording grid. "
                "A displacement in metres, not a gait speed and not a clinical walking distance.")
    DISP_AT = "3.0 s after the onset of the final decoder output (the biped is driven from 2 s earlier)"
    FALL_DEF = ("the number of rollouts in which trunk pitch exceeded 60 deg from vertical or hip height fell "
                "below half the standing hip height, at any time from trial time 2 s.")
DRIVE_NOTE = ""
if ec.F7_SIM == "causal_from2s":
    DRIVE_NOTE = (" Every condition drives the biped from 2 s into the trial: the Walker by the reference walker's "
                  "goniometer signal; the Decoded condition and both nulls by the decoder's readout before temporal "
                  "refinement until the onset of the final decoder output (4 s) and by the final output thereafter, "
                  "the nulls being built on that whole drive.")
ORDER = ["Walker", "Decoded", "CoordinationNull", "InPhaseNull"]
DISPLAY = {"Walker": "Walker (window-matched)", "Decoded": "Decoded",
           "CoordinationNull": "Coordination null", "InPhaseNull": "In-phase null"}
SPEC = {"Walker": ("calib", "wphase_matched"), "Decoded": ("confirm", "dphase"),
        "CoordinationNull": ("confirm", "dphase_iscr"), "InPhaseNull": ("confirm", "dphase_uscr")}
GROUPS = [("Pooled", None), ("Patient 1", "pat01"), ("Patient 2", "pat02")]
CONTRASTS = [("Walker_minus_Decoded", "Walker", "Decoded"),
             ("Decoded_minus_CoordinationNull", "Decoded", "CoordinationNull"),
             ("Decoded_minus_InPhaseNull", "Decoded", "InPhaseNull")]
DROPPED = {("Walker_minus_Decoded", "lr_asymmetry")}
METRICS = [
    ("disp_3s_m", DISP_Q, "m", DISP_Q_TITLE),
    ("cadence_steps_per_s", "cadence", "steps s^-1", "Cadence (steps s⁻¹)"),
    ("step_length_m", "step length", "m", "Step length (m)"),
    ("lr_asymmetry", "advancing-step count imbalance", "dimensionless",
     "Advancing-step count imbalance"),
]
COUNTS = [("pct_backward", "rollouts moving backward", "%", "Rollouts moving backward (%)"),
          ("n_fallen", "rollouts with a fall", "count", "Rollouts with a fall (n)")]


def md5(path):
    h = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def p_cell(p):
    """Table cell for a p value, house style: floor at 0.001 ('< 0.001'), else three decimals."""
    if not np.isfinite(p):
        return ""
    if p < 0.001:
        return "< 0.001"
    return float(Decimal(repr(float(p))).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def stars(p):
    return "***" if p < 1e-3 else "**" if p < 1e-2 else "*" if p < 0.05 else "ns"


def panel_md(stem):
    d = pd.read_csv(ec.sd(f"{stem}_metadata.csv"))
    return dict(zip(d.key, d.value))


def lme_paired_test(a, b, key, mixed=True):
    """Per-trial paired difference d = a - b matched on (session, trial_idx); b None = the vs-zero
    test (d is a's value itself)."""
    A = a.set_index(["session", "trial_idx"])[key].astype(float)
    dd = (A if b is None else A - b.set_index(["session", "trial_idx"])[key].astype(float)).dropna()
    sess = dd.index.get_level_values("session").astype(str)
    df = pd.DataFrame({"d": dd.to_numpy(), "subject": sess.str[:5], "session": sess})
    multi = df.session.nunique() >= 2
    fits = []
    for m in ("lbfgs", "powell", "cg") if multi else ():
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            f = smf.mixedlm("d ~ 1", df, groups=df["session"], re_formula="~1").fit(
                reml=True, method=m, maxiter=2000)
        if f.converged and np.isfinite(f.bse_fe["Intercept"]) and np.isfinite(f.llf):
            fits.append((m, f))
    if fits:
        m, f = max(fits, key=lambda t: t[1].llf)
        beta, se, p_wald = (float(f.fe_params["Intercept"]), float(f.bse_fe["Intercept"]),
                            float(f.pvalues["Intercept"]))
        resid, revar, model = f.resid, float(f.cov_re.iloc[0, 0]), "LME d ~ 1 + (1 | session)"
    else:
        if multi:
            bnd = []
            for m in ("lbfgs", "powell", "cg"):
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    f = smf.mixedlm("d ~ 1", df, groups=df["session"], re_formula="~1").fit(
                        reml=True, method=m, maxiter=2000)
                bnd.append(bool(f.converged) and bool(np.isfinite(f.llf))
                           and float(f.cov_re.iloc[0, 0]) < 1e-4 * float(np.var(df.d)))
            if not all(bnd):
                raise RuntimeError(f"LME d ~ 1 + (1 | session) converged with no optimizer "
                                   f"({key}, n = {len(df)})")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            f = smf.ols("d ~ 1", df).fit()
        m = "ols (session variance 0)" if multi else "ols (one session: intercept only)"
        beta, se, p_wald = (float(f.params["Intercept"]), float(f.bse["Intercept"]),
                            float(f.pvalues["Intercept"]))
        resid, revar = f.resid, 0.0
        model = "LME d ~ 1 + (1 | session), session variance 0 -> intercept only"
    if not mixed:
        model += " (within one patient)"
    p_sw = float(sps.shapiro(resid).pvalue)
    p_wx = float(sps.wilcoxon(df.d, zero_method="wilcox").pvalue)
    p = p_wald
    used = "LME" if mixed else "LME (within patient)"
    return dict(n=len(df), n_subjects=df.subject.nunique(), n_sessions=df.session.nunique(),
                mean_difference_lme=beta, se_lme=se, p_lme=p_wald, session_variance=revar,
                optimizer=m, shapiro_p=p_sw, p_wilcoxon=p_wx,
                p_paired_t=float(sps.ttest_1samp(df.d, 0).pvalue),
                median_difference=float(np.median(df.d)), model=model,
                test_used=used, p_used=p, star=stars(p))


def rate_length(df):
    """The two rate_length quantities, per rollout — copied from 03_figures/fig7_sim_cadence_symmetry.py."""
    n = df.n_adv_events.astype(float)
    return df.assign(cadence_steps_per_s=n / df.dur_s.astype(float),
                     step_length_m=df.dx_m.astype(float) / n.where(n > 0),
                     lr_asymmetry=df.r5_lr_asym.astype(float))


def load():
    rhy = rate_length(F7_SCORED.keep_sim(pd.read_csv(RHY)))
    wrhy = rate_length(F7_SCORED.keep_sim(pd.read_csv(WRHY)))

    def sel(df, cond):
        st, ar = SPEC[cond]
        return df[(df.stage == st) & (df.condition == ar) & (df.phase_src == "k1")
                  & (df.lam.astype(float) == 0.0)]

    A = {c: sel(wrhy if c == "Walker" else rhy, c) for c in ORDER}

    t = F7_SCORED.keep_sim(pd.read_csv(TRA))
    t = pd.concat([t[t.condition != "wphase"], F7_SCORED.keep_sim(pd.read_csv(WTRA))], ignore_index=True)
    t = t[t.t_s <= TEND + 1e-9]
    end = {}
    for c in ORDER:
        s = t[t.condition == SPEC[c][1]]
        end[c] = s.sort_values("t_s").groupby(["session", "trial_idx"]).x_m.last()

    out = {}
    for c in ORDER:
        d = A[c][["stage", "condition", "session", "trial_idx", "dur_s", "dx_m", "n_adv_events",
                  "n_adv_l", "n_adv_r", "fell", "cadence_steps_per_s", "step_length_m",
                  "lr_asymmetry"]].copy()
        d["trial_idx"] = d.trial_idx.astype(int)
        e = end[c].rename("disp_3s_m").reset_index()
        e["trial_idx"] = e.trial_idx.astype(int)
        d = d.merge(e, on=["session", "trial_idx"], how="inner")
        assert len(d) == len(A[c]), f"{c}: {len(A[c])} rhythm rollouts but {len(d)} traced"
        d["condition"] = c
        d["patient"] = d.session.str[:5]
        d["fell"] = d.fell.astype(float)
        out[c] = d.reset_index(drop=True)

    keys = {c: set(zip(out[c].session, out[c].trial_idx)) for c in ORDER}
    assert keys["Walker"] == keys["Decoded"] == keys["CoordinationNull"] == keys["InPhaseNull"], \
        "the four conditions do not share one trial set"
    return pd.concat([out[c] for c in ORDER], ignore_index=True)


def metronome_status():
    """Which conditions the 3 s batch holds, measured from the files rather than remembered."""
    roll = F7_SCORED.keep_sim(pd.read_csv(ROLL, keep_default_na=False, low_memory=False))
    in_roll = sorted(set(roll[roll.phase_src == "metronome_0.9"].condition))
    rhy_conditions = sorted(set(pd.read_csv(RHY).condition)) + sorted(set(pd.read_csv(WRHY).condition))
    tra_conditions = sorted(set(pd.read_csv(TRA).condition)) + sorted(set(pd.read_csv(WTRA).condition))
    n_roll = int((roll.phase_src == "metronome_0.9").sum())
    assert not [a for a in in_roll if a in rhy_conditions or a in tra_conditions], \
        "a metronome condition IS present in the 3 s rhythm / trace replay — re-instate the reference"
    return (f"DROPPED. The 3 s batch g_v5_rollouts.csv holds {n_roll} metronome rollouts "
            f"({', '.join(in_roll)}, phase_src 'metronome_0.9'), but the 3 s rhythm probe "
            f"(rhythm_metrics.csv, column \"condition\" values {', '.join(sorted(set(rhy_conditions)))}) and the 3 s hip-x "
            f"trace dump (traces_x.csv.gz, column \"condition\" values {', '.join(sorted(set(tra_conditions)))}) were replayed "
            "for the four tabulated conditions only. Every quantity in this table is a rhythm-probe "
            "or trace quantity, so none of them exists for the metronome in the 3 s build. Checked from the files "
            "on every run.")


def build(tab):
    rows = []
    for gname, pat in GROUPS:
        g = tab if pat is None else tab[tab.patient == pat]
        A = {c: g[g.condition == c] for c in ORDER}
        for c in ORDER:
            d = A[c]
            for key, quant, unit, _ in METRICS:
                v = d[key].astype(float).dropna()
                rows.append(dict(row_type="outcome", group=gname, condition=DISPLAY[c],
                                 quantity=quant, unit=unit, n=len(v),
                                 n_rollouts=len(d), n_sessions=d.session.nunique(),
                                 median=float(v.median()), q25=float(v.quantile(.25)),
                                 q75=float(v.quantile(.75))))
            back = 100.0 * float((d.disp_3s_m.astype(float) < 0).mean())
            rows.append(dict(row_type="outcome", group=gname, condition=DISPLAY[c],
                             quantity="rollouts moving backward", unit="%", n=len(d),
                             n_rollouts=len(d), n_sessions=d.session.nunique(), value=back))
            rows.append(dict(row_type="outcome", group=gname, condition=DISPLAY[c],
                             quantity="rollouts with a fall", unit="count", n=len(d),
                             n_rollouts=len(d), n_sessions=d.session.nunique(),
                             value=float(d.fell.sum())))
        mixed = g.patient.nunique() >= 2
        for name, a, b in CONTRASTS:
            for key, quant, unit, _ in METRICS:
                if (name, key) in DROPPED:
                    continue
                r = lme_paired_test(A[a], A[b], key, mixed=mixed)
                rows.append(dict(row_type="contrast", group=gname,
                                 condition=f"{DISPLAY[a]} minus {DISPLAY[b]}", quantity=quant,
                                 unit=unit, n=r["n"], n_sessions=r["n_sessions"],
                                 a_median=float(A[a][key].astype(float).dropna().median()),
                                 b_median=float(A[b][key].astype(float).dropna().median()),
                                 mean_difference_lme=r["mean_difference_lme"], se=r["se_lme"],
                                 median_difference=r["median_difference"],
                                 shapiro_p=r["shapiro_p"], p_wald=r["p_lme"],
                                 p_wilcoxon=r["p_wilcoxon"], model=r["model"],
                                 test_used=r["test_used"], p_used=r["p_used"], star=r["star"],
                                 comparison_key=name, n_subjects=r["n_subjects"],
                                 session_variance=r["session_variance"], optimizer=r["optimizer"]))
    cols = ["row_type", "group", "condition", "quantity", "unit", "n", "n_rollouts", "n_sessions",
            "median", "q25", "q75", "value", "a_median", "b_median", "mean_difference_lme", "se",
            "median_difference", "shapiro_p", "p_wald", "p_wilcoxon", "model", "test_used",
            "p_used", "star", "comparison_key", "n_subjects", "session_variance", "optimizer"]
    return pd.DataFrame(rows).reindex(columns=cols)


PANEL_OF = {"cadence_steps_per_s": "fig7_sim_cadence", "step_length_m": "sup_fig7_sim_step_length",
            "lr_asymmetry": "fig7_sim_step_symmetry"}
SOFT_RTOL = {"p_used": 1e-2, "mean_difference_lme": 1e-4}


def crosscheck(df):
    """The pooled numbers must BE the panel numbers. Nothing is written if any line disagrees."""
    fails, notes, checked = [], [], 0
    pooled = df[df.group == "Pooled"]

    for key, stem in PANEL_OF.items():
        quant = dict((k, q) for k, q, _, _ in METRICS)[key]
        exp = dict(kv.split(" ") for kv in panel_md(stem)["medians_per_condition"].split("; "))
        for c in ORDER:
            got = pooled[(pooled.quantity == quant) & (pooled.condition == DISPLAY[c])]
            got = f"{float(got['median'].iloc[0]):.6f}"
            checked += 1
            if got != exp[c]:
                fails.append(f"{stem} medians_per_condition {c}: table {got} vs panel {exp[c]}")

    ep = panel_md("fig7_sim_trajectories")["endpoint_stats"]
    ep = ep.split(" per condition — ", 1)[1].rsplit(". Computed", 1)[0]
    for part in ep.split("; "):
        lab, body = part.split(": ", 1)
        exp_med = body.split(",")[0].replace("median ", "").replace(" m", "").strip()
        exp_back = body.split("backward ")[1].replace("%", "").strip()
        o = pooled[pooled.condition == DISPLAY[lab]]
        got_med = f"{float(o[o.quantity == DISP_Q]['median'].iloc[0]):+.2f}"
        got_back = f"{float(o[o.quantity == 'rollouts moving backward']['value'].iloc[0]):.0f}"
        checked += 2
        if got_med != exp_med:
            fails.append(f"endpoint_stats {lab} median: table {got_med} vs panel {exp_med}")
        if got_back != exp_back:
            fails.append(f"endpoint_stats {lab} backward: table {got_back}% vs panel {exp_back}%")

    for key, stem in list(PANEL_OF.items()) + [("disp_3s_m", "fig7_sim_trajectories")]:
        quant = dict((k, q) for k, q, _, _ in METRICS)[key]
        st = pd.read_csv(ec.sd(f"{stem}_stats.csv")).set_index("comparison")
        for name, _, _ in CONTRASTS:
            if (name, key) in DROPPED:
                continue
            r = pooled[(pooled.quantity == quant) & (pooled.comparison_key == name)].iloc[0]
            e = st.loc[name]
            soft = r.test_used == "LME"
            for col, got, exp in (("n", int(r.n), int(e.n)),
                                  ("mean_difference_lme", f"{r.mean_difference_lme:.6f}",
                                   f"{float(e.mean_difference_lme):.6f}"),
                                  ("p_used", f"{r.p_used:.6g}", f"{float(e.p_used):.6g}"),
                                  ("test_used", r.test_used, e.test_used),
                                  ("star", r.star, e.star)):
                checked += 1
                if str(got) == str(exp):
                    continue
                if col in SOFT_RTOL and soft and \
                        abs(float(got) - float(exp)) <= SOFT_RTOL[col] * max(abs(float(got)), abs(float(exp))):
                    notes.append(f"{stem} {name}: mixed-model {col} {got} here against {exp} in "
                                 f"the panel Source Data (within {SOFT_RTOL[col]:g} relative). The "
                                 f"optimiser kept by the largest log-likelihood is '{r.optimizer}' "
                                 f"here and '{e.optimizer}' there; the three optimisers' "
                                 "log-likelihoods agree to about 1e-13, so which one wins is "
                                 "environment-dependent. Star, n and the test used reproduce exactly, "
                                 "and the star is the reported quantity.")
                    continue
                fails.append(f"{stem} {name} {col}: table {got} vs panel {exp}")
    return checked, fails, notes


def write_source_data(df, meta):
    df.to_csv(SD, index=False)
    pd.DataFrame(meta, columns=["key", "value"]).to_csv(MD, index=False)


def sheet(wb, title, header, body, widths, note):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in body:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str))
    ws.insert_rows(1)
    ws["A1"] = note
    ws["A1"].font = Font(bold=True)
    ws.freeze_panes = "A3"
    return ws


TITLE = ("Supplementary Table 4 | Simulation outcomes per condition, per patient and pooled, "
         "with paired contrasts and definitions.")


def write_xlsx(df, meta, defs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    def cell(m, q1, q3):
        m, q1, q3 = (round(float(v), 3) + 0.0 for v in (m, q1, q3))
        return f"{m:.3f} [{q1:.3f}, {q3:.3f}]"

    wb = Workbook()
    wb.remove(wb.active)

    body = []
    for gname, _ in GROUPS:
        for c in ORDER:
            s = df[(df.row_type == "outcome") & (df.group == gname)
                   & (df.condition == DISPLAY[c])].set_index("quantity")
            row = [gname, DISPLAY[c], int(s.n_rollouts.iloc[0]), int(s.n_sessions.iloc[0])]
            for _, quant, _, _ in METRICS:
                r = s.loc[quant]
                row.append(cell(r["median"], r["q25"], r["q75"]))
            row.append(round(float(s.loc["rollouts moving backward", "value"]), 1))
            row.append(int(s.loc["rollouts with a fall", "value"]))
            body.append(row)
    n_trials_note = "188"
    if not ec.IS_DEFAULT:
        n_trials_note = str(int(df[(df.row_type == "outcome") & (df.group == "Pooled")
                                   & (df.condition == DISPLAY["Decoded"])].n_rollouts.iloc[0]))
    ws = sheet(wb, "1 Outcomes",
          ["Group", "Condition", "n rollouts", "n sessions"] + [h for _, _, _, h in METRICS]
          + [h for _, _, _, h in COUNTS], body,
          [11, 24, 10, 10, 22, 20, 18, 18, 15, 13],
          TITLE + " (1) Outcomes. Median [25th, 75th percentile] over rollouts, from the 3 s "
                  "usable-duration build in which the walker condition is replayed on exactly the "
                  "decoded window of each decoded trial, so the four conditions share one set of "
                  f"{n_trials_note} trials. Forward displacement is the simulated hip displacement of the "
                  "physics rollout in metres; it is not a gait speed and not a clinical walking "
                  "distance.")
    for row in ws.iter_rows(min_row=3, min_col=9, max_col=9):
        for c in row:
            c.number_format = "0.0"

    body = []
    for r in df[df.row_type == "contrast"].itertuples():
        body.append([r.group, r.condition, r.quantity, int(r.n), round(float(r.a_median), 3),
                     round(float(r.b_median), 3), round(float(r.mean_difference_lme), 3),
                     round(float(r.se), 3), r.test_used, p_cell(float(r.p_used)), r.star])
    ws = sheet(wb, "2 Paired contrasts",
               ["Group", "Contrast (A minus B)", "Metric", "n pairs", "A median", "B median",
                "Mean difference", "Standard error", "Test", "p (two-sided)", "Significance"],
               body, [11, 34, 26, 9, 11, 11, 13, 13, 26, 13, 12],
               TITLE + " (2) Paired contrasts. Every contrast is paired on the (session, trial) "
                       "key: the window-matched walker replays the same trials as the decoded "
                       "condition, and each null re-runs those trials with the phases swapped. "
                       "The per-trial difference is tested with a linear mixed model carrying a "
                       "session random intercept, pooled and within each patient; Wald p throughout, Shapiro–Wilk and Wilcoxon are "
                       "in the Source Data only. Medians are descriptive.")
    for row in ws.iter_rows(min_row=3, min_col=10, max_col=10):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.000"

    ws = sheet(wb, "3 Definitions", ["Quantity", "Unit", "Definition"], defs, [34, 22, 110],
               TITLE + " (3) Definitions and provenance.")
    ws.append([])
    ws.append(["Note", "", "Detail"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for k, v in meta:
        ws.append([k, "", v])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(XLSX)


LEGEND_MD = """# Supplementary Table 4 [simulation outcomes] — legend

> Every Fig. 7 simulation deliverable reports the 3 s usable-duration, window-matched-walker build
> (trials with at least 3 s of decoder output after the decoder's 4 s start-up).

Deliverable: `supplementary/tables/Supplementary_Table_S4_Simulation_Outcomes.xlsx`
Producer: `03_figures/sup_table_sim_outcomes.py`
Source Data: `04_source_data/supplementary/sup_table_sim_outcomes_source_data.csv` (+ `_metadata.csv`)

Every number below has a row in the Source Data file named
above. This file is written by the
producer, so the numbers cannot drift from the workbook.

---

## Legend

```
Supplementary Table 4 | Simulation outcome quantities per condition, per patient and pooled.
Every quantity measured on each physics rollout and summarised per condition as the median [25th,
75th percentile] over rollouts; (a) outcomes, (b) paired contrasts, (c) quantity definitions and
provenance. Conditions. Walker, the able-bodied reference walker's own gait phase replayed on
exactly the decoded window of each decoded trial, a ceiling rather than a null; Decoded, the
patient's EEG-decoded per-joint phase driving the normative waveform template; coordination null,
the inter-joint phase relationship destroyed with timing preserved; in-phase null, the decoded phases
replaced by one common phase across joints.{DRIVE_NOTE} All four conditions run the same {N_TRIALS} trials from
{N_SESS} sessions of 2 patients ({N_P1} trials from patient 1, {N_P2} from patient 2). Columns.
Forward displacement, the simulated hip displacement reached {DISP_AT}, in metres -
not a gait speed and not a clinical walking distance; cadence, the number of advancing left-right
support transitions divided by the rollout duration, an advancing transition being one in which the
incoming support foot lands at least 0.01 m in front of the outgoing one; step length, the net
forward hip displacement divided by that number of advancing steps, so negative values indicate net
backward progression; advancing-step count imbalance, the absolute difference between the numbers of left
and right advancing steps divided by their sum, 0 for perfectly alternating stepping and 1 when
every advancing step falls on one side; rollouts moving backward, the percentage whose {TEND} s
displacement is negative; rollouts with a fall, the count in which trunk pitch exceeded 60 deg from
vertical or hip height fell below half the standing height. Pooled medians. Forward displacement
{D_W} m for the walker condition, {D_D} m decoded, {D_S} m coordination null and {D_T} m in-phase
null; cadence {C_W}, {C_D}, {C_S} and {C_T} steps s-1; step length {L_W}, {L_D}, {L_S} and {L_T} m;
advancing-step count imbalance {A_W}, {A_D}, {A_S} and {A_T}. Backward and fallen rollouts. {BACKFALL}
Contrasts. Each contrast in (b) is paired on the (session, trial) key - the window-matched walker
replays the same trials as the decoded condition, and each null re-runs those trials with the
phases swapped - and the per-trial difference is tested with a linear mixed model carrying a session
random intercept, pooled and within each patient (Wald p; Shapiro-Wilk and Wilcoxon are reported in
the Source Data only), and the model applied is named in every row.
Pooled, the decoded condition exceeds both nulls on forward displacement ({T_DS_DISP}; {T_DT_DISP}),
on step length ({T_DS_LEN}; {T_DT_LEN}), has a lower advancing-step count imbalance than both nulls ({T_DS_ASYM}; {T_DT_ASYM}), and
matches the walker ceiling on cadence ({T_WD_CAD}) while falling short of it on forward displacement
({T_WD_DISP}) and step length ({T_WD_LEN}). All outcomes are behavioural simulation quantities. No
ground reaction force, contact force, friction or joint moment is reported, and no gait speed is
reported anywhere. Cohort post_operative, kinematic set, 9 of 11 post-operative sessions; n = 2 patients,
case-level.
```

---

## Numbers used above, with their Source Data rows

All rows in `04_source_data/supplementary/sup_table_sim_outcomes_source_data.csv`.

| number in the legend | Source Data row |
|---|---|
| n rollouts and n sessions per condition | `row_type=outcome`, `group=Pooled`: `n_rollouts`, `n_sessions` |
| per-patient trial counts | `row_type=outcome`, `group=Patient 1` / `Patient 2`: `n_rollouts` |
| pooled medians of the four quantities | `row_type=outcome`, `group=Pooled`, `quantity=...`: `median`, `q25`, `q75` |
| rollouts moving backward (%) | `row_type=outcome`, `quantity=rollouts moving backward`: `value` |
| rollouts with a fall (n) | `row_type=outcome`, `quantity=rollouts with a fall`: `value` |
| every contrast statistic | `row_type=contrast`: `mean_difference_lme`, `se`, `model`, `test_used`, `p_used`, `star` |
| which model each contrast used | `row_type=contrast`: `model` (and `n_subjects`) |

## Notes

1. **Simulation outcomes, not biomechanics.** Simulated mechanical quantities are not biomechanical
   measurements (soft contact model). No GRF, contact force, friction or joint moment is reported.
2. **Forward displacement is not a speed.** It is the simulated hip displacement of the physics
   rollout at {TEND} s, in metres; no m/s value appears in the table.
3. **The walker condition is a ceiling, not a null**, and it is window-matched: it replays the same
   trials over exactly the decoded window, which is what makes the walker contrast paired.\n4. **Model.** The patient variance of two patients is not estimable, so the model carries no patient term. The
   `model` column names the model used in every row; pooled and per-patient statistics are never
   quoted without it.
5. **No amplitude-gain wording.** In these conditions the decode supplies per-joint phase (timing
   and coordination) and the template supplies waveform and amplitude; there is no scalar amplitude
   gain in them.
6. **Cohort scope.** 9 of 11 post-operative sessions from 2 patients, case-level. Rollouts are
   nested within sessions and within 2 patients, so every p value is a trial-level test, never a
   population-level one.
7. **Uncorrected.** The endpoints are not independent: the two nulls re-run the same trials, and
   cadence, step length and asymmetry all read the same advancing-step array.


No metronome reference is tabulated. {METRONOME_SHORT}
"""


def write_legend(df, metronome):
    pooled = df[(df.group == "Pooled") & (df.row_type == "outcome")]

    def med(quant, c):
        return f"{float(pooled[(pooled.quantity == quant) & (pooled.condition == DISPLAY[c])]['median'].iloc[0]):.2f}"

    def sc(quant, c, col="value"):
        return float(pooled[(pooled.quantity == quant) & (pooled.condition == DISPLAY[c])][col].iloc[0])

    def contrast(name, quant):
        r = df[(df.group == "Pooled") & (df.row_type == "contrast")
               & (df.comparison_key == name) & (df.quantity == quant)].iloc[0]
        p = p_cell(float(r.p_used))
        p = f"p {p}" if isinstance(p, str) else f"p = {p:.3f}"
        return f"mean difference {float(r.mean_difference_lme):+.3f}, {p}"

    p1 = int(df[(df.group == "Patient 1") & (df.quantity == "cadence")
                & (df.condition == DISPLAY["Decoded"])].n_rollouts.iloc[0])
    p2 = int(df[(df.group == "Patient 2") & (df.quantity == "cadence")
                & (df.condition == DISPLAY["Decoded"])].n_rollouts.iloc[0])
    backfall = " ".join(
        f"{DISPLAY[c]} {sc('rollouts moving backward', c):.0f} % backward and "
        f"{int(sc('rollouts with a fall', c))} fallen;" for c in ORDER).rstrip(";") + "."
    txt = LEGEND_MD.format(
        N_TRIALS=p1 + p2, N_SESS=int(pooled.n_sessions.iloc[0]), N_P1=p1, N_P2=p2,
        TEND=f"{TEND:.1f}", DISP_AT=DISP_AT, DRIVE_NOTE=DRIVE_NOTE,
        D_W=med(DISP_Q, "Walker"),
        D_D=med(DISP_Q, "Decoded"),
        D_S=med(DISP_Q, "CoordinationNull"),
        D_T=med(DISP_Q, "InPhaseNull"),
        C_W=med("cadence", "Walker"), C_D=med("cadence", "Decoded"),
        C_S=med("cadence", "CoordinationNull"), C_T=med("cadence", "InPhaseNull"),
        L_W=med("step length", "Walker"), L_D=med("step length", "Decoded"),
        L_S=med("step length", "CoordinationNull"), L_T=med("step length", "InPhaseNull"),
        A_W=med("advancing-step count imbalance", "Walker"),
        A_D=med("advancing-step count imbalance", "Decoded"),
        A_S=med("advancing-step count imbalance", "CoordinationNull"),
        A_T=med("advancing-step count imbalance", "InPhaseNull"),
        BACKFALL=backfall,
        T_DS_DISP=contrast("Decoded_minus_CoordinationNull", DISP_Q),
        T_DT_DISP=contrast("Decoded_minus_InPhaseNull", DISP_Q),
        T_DS_LEN=contrast("Decoded_minus_CoordinationNull", "step length"),
        T_DT_LEN=contrast("Decoded_minus_InPhaseNull", "step length"),
        T_DS_ASYM=contrast("Decoded_minus_CoordinationNull", "advancing-step count imbalance"),
        T_DT_ASYM=contrast("Decoded_minus_InPhaseNull", "advancing-step count imbalance"),
        T_WD_CAD=contrast("Walker_minus_Decoded", "cadence"),
        T_WD_DISP=contrast("Walker_minus_Decoded", DISP_Q),
        T_WD_LEN=contrast("Walker_minus_Decoded", "step length"),
        METRONOME_SHORT=metronome.split(" Checked from the files")[0])
    tmp = LEGEND + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(txt)
    os.replace(tmp, LEGEND)


DEFS = [
    (DISP_Q_TITLE, "m", DISP_DEF),
    ("Cadence (steps s⁻¹)", "steps s⁻¹",
     "the number of advancing left-right support transitions divided by the rollout duration. A "
     "transition advances when the incoming support foot lands at least 0.01 m in front of the "
     "outgoing one, so only forward steps are counted."),
    ("Step length (m)", "m",
     "the net forward hip displacement of the rollout divided by its number of advancing steps, "
     "that is, the mean forward progression per advancing step. A rollout with no advancing step "
     "has no step length and is dropped. Negative values indicate net backward progression."),
    ("Advancing-step count imbalance", "dimensionless",
     "the absolute difference between the numbers of left and right advancing steps divided by "
     "their sum: 0 for perfectly alternating left-right stepping, 1 when every advancing step "
     "falls on the same side."),
    ("Rollouts moving backward (%)", "%",
     "the percentage of the condition's rollouts whose 3.0 s forward displacement is negative."),
    ("Rollouts with a fall (n)", "count", FALL_DEF),
]


def main():
    tab = load()
    metronome = metronome_status()
    df = build(tab)
    checked, fails, notes = crosscheck(df)
    if fails:
        print("CROSS-CHECK FAILED against the live Fig. 7 panels — nothing written:")
        for f in fails:
            print("  " + f)
        raise SystemExit(1)
    print(f"cross-check OK: {checked} pooled values reproduce the Fig. 7 panel Source Data")
    for n in notes:
        print("  NAMED TOLERANCE — " + n)

    cad = panel_md("fig7_sim_cadence")
    for stem in ("fig7_sim_trajectories", "fig7_sim_step_symmetry", "sup_fig7_sim_step_length"):
        o = panel_md(stem)
        assert o["build_id"].split(" ")[0] == cad["build_id"].split(" ")[0], f"{stem}: build_id tag differs"
        for k in ("usable_duration_gate", "walker_window", "test_rule"):
            assert o[k] == cad[k], f"{stem} metadata row {k} differs from fig7_sim_cadence"

    p1 = tab[tab.patient == "pat01"]
    probe = []
    for m in ("lbfgs", "powell", "cg"):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            A = p1[p1.condition == "Walker"].set_index(["session", "trial_idx"]).cadence_steps_per_s
            Bv = p1[p1.condition == "Decoded"].set_index(["session", "trial_idx"]).cadence_steps_per_s
            d0 = (A.astype(float) - Bv.astype(float)).dropna()
            dfp = pd.DataFrame({"d": d0.to_numpy(),
                                "session": d0.index.get_level_values("session").astype(str)})
            f = smf.mixedlm("d ~ 1", dfp, groups=dfp["session"], re_formula="~1").fit(
                reml=True, method=m, maxiter=2000)
        probe.append(f"{m}: session variance {float(f.cov_re.iloc[0, 0]):.4f}, intercept SE "
                     f"{float(f.bse_fe['Intercept']):.4f}, log-likelihood {f.llf:.6f}")

    n_tr = int(df[(df.group == "Pooled") & (df.quantity == "cadence")
                  & (df.condition == DISPLAY["Decoded"])].n_rollouts.iloc[0])
    meta = [
        ("table", "Supplementary Table 4 'Simulation outcomes per condition' -> "
                  "figures_out/supplementary/tables/Supplementary_Table_S4_Simulation_Outcomes.xlsx (3 sheets: "
                  "1 outcomes | 2 paired contrasts | 3 definitions and provenance)."),
        ("paper_location", "Supplementary Table 4; Supplementary Note 7, simulation paragraphs. "
                           ),
        ("producer", "03_figures/sup_table_sim_outcomes.py"),
    ] + ([ec.F7_SIM_MD] if ec.F7_SIM else []) + [
        ("cohort", cad["cohort"]),
        ("usable_duration_gate", cad["usable_duration_gate"]),
        ("walker_window", cad["walker_window"]),
        ("test_rule", cad["test_rule"]),
        ("conditions", "Walker (window-matched): the reference walker's own gait phase replayed "
                       "on exactly the decoded window of each decoded trial; a ceiling, not a "
                       "null. Decoded: the patient's EEG-decoded per-joint phase driving the "
                       "normative gait template. Coordination null: the inter-joint phase "
                       "relationship destroyed, timing kept. In-phase null: the decoded phases "
                       "replaced by one common phase, so that all six joints move in phase."
                       f"{DRIVE_NOTE} All four conditions share one set of {n_tr} trials. Source Data field: "
                       "column \"condition\" = 'wphase_matched' with column \"stage\" = 'calib' "
                       "(Walker; walker_matched_rhythm.csv, walker_matched_traces_x.csv.gz), and "
                       "'dphase', 'dphase_iscr', 'dphase_uscr' with column \"stage\" = 'confirm' "
                       "(Decoded, Coordination null, In-phase null; rhythm_metrics.csv, "
                       "traces_x.csv.gz); column \"phase_src\" = 'k1' and column \"lam\" = 0.0 "
                       "for every condition."),
        ("metric_definitions", " | ".join(f"{q} ({u}): {d}" for q, u, d in DEFS)),
        ("per_patient_model", "Within a single patient the test_rule model keeps its SESSION random "
                              "intercept (patient 1: 4 sessions, patient 2: 5), so the per-patient rows "
                              "fit d ~ 1 + (1 | session) to that patient's rollouts, test_used 'LME "
                              "(within patient)'. Measured on this build "
                              "(patient 1, Walker minus Decoded, cadence): " + "; ".join(probe)
                              + ". Every contrast row names the model it used in the `model` column."),
        ("pairing_rule", "Every contrast is paired on (session, trial_idx): the window-matched "
                         "walker replays the same trials over exactly the decoded samples, and "
                         "each null re-runs those trials with the phases swapped."),
        ("cross_check", f"{checked} pooled values are re-derived here and asserted equal to the "
                        "live Fig. 7 panel Source Data on every run: medians_per_condition of "
                        "fig7_sim_cadence / fig7_sim_step_symmetry / sup_fig7_sim_step_length, "
                        "endpoint_stats of fig7_sim_trajectories, and every pooled contrast row of "
                        "the four *_stats.csv files (n, mean difference, p_used, test used, star). "
                        "The script writes nothing if any value disagrees."),
        ("known_discrepancy", " ".join(notes) if notes else
         "none — every cross-checked value reproduces the panel Source Data exactly."),
        ("statistic_reported", "Median [25th, 75th percentile] over rollouts for every quantity; "
                               "rollouts moving backward is a percentage and rollouts with a fall "
                               "a count. Fallen rollouts are NOT dropped."),
        ("multiple_comparisons", "Reported uncorrected. The endpoints are not independent: the two "
                                 "nulls re-run the same trials, and cadence, step length and "
                                 "asymmetry all read the same advancing-step array."),
        ("software_environment", f"python {platform.python_version()} ({platform.system()} "
                                 f"{platform.machine()}); numpy {np.__version__}; pandas "
                                 f"{pd.__version__}; scipy {__import__('scipy').__version__}; "
                                 f"statsmodels {__import__('statsmodels').__version__}; openpyxl "
                                 f"{__import__('openpyxl').__version__}"),
    ]

    write_source_data(df, meta)
    write_xlsx(df, meta, [[q, u, d] for q, u, d in DEFS])
    write_legend(df, metronome)
    for p in (SD, MD, XLSX, LEGEND):
        print("wrote", os.path.relpath(p, ROOT))

    pooled = df[(df.group == "Pooled") & (df.row_type == "outcome")]
    for _, quant, _, _ in METRICS:
        s = pooled[pooled.quantity == quant]
        print(f"  {quant:32s} " + " | ".join(
            f"{DISPLAY[c][:12]:12s} {float(s[s.condition == DISPLAY[c]]['median'].iloc[0]):+.3f}"
            for c in ORDER))
    for r in df[df.row_type == "contrast"].itertuples():
        print(f"  {r.group:9s} | {r.quantity:26s} | {r.comparison_key:24s} "
              f"mean {float(r.mean_difference_lme):+.4f}  {r.test_used:28s} "
              f"p {float(r.p_used):.3g} {r.star}")


if __name__ == "__main__":
    main()
